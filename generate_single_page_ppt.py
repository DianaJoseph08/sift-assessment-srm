import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_single_page_presentation(output_path):
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # ----------------------------------------------------
    # COLOR PALETTE
    # ----------------------------------------------------
    BG_CANVAS        = RGBColor(241, 245, 249)  # Slate-100 background
    NAVY_HERO        = RGBColor(11, 23, 44)     # Deep Slate-950
    MOTHERSON_BLUE   = RGBColor(3, 77, 161)     # Primary brand blue
    MOTHERSON_RED    = RGBColor(225, 29, 72)    # Brand accent red
    CYAN_ACCENT      = RGBColor(14, 165, 233)   # Sky-500
    EMERALD_GREEN    = RGBColor(16, 185, 129)   # Emerald-500
    AMBER_GOLD       = RGBColor(245, 158, 11)   # Amber-500
    PURPLE_ACCENT    = RGBColor(168, 85, 247)   # Vibrant Purple
    PURPLE_HEADER    = RGBColor(192, 132, 252)  # High-contrast light purple for dark header
    ROSE_ACCENT      = RGBColor(244, 63, 94)    # Rose-500
    CARD_BG          = RGBColor(255, 255, 255)  # Pure White
    CARD_BORDER      = RGBColor(203, 213, 225)  # Slate-300
    PILL_BG          = RGBColor(23, 37, 63)     # Deep Blue-Slate
    TEXT_DARK        = RGBColor(15, 23, 42)     # Slate-900
    TEXT_MUTED       = RGBColor(100, 116, 139)  # Slate-500
    TEXT_WHITE       = RGBColor(255, 255, 255)  # White
    ROW_BG_LIGHT     = RGBColor(248, 250, 252)  # Slate-50

    # 1. Slide Canvas Background
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BG_CANVAS
    bg_shape.line.fill.background()

    # ----------------------------------------------------
    # HEADER BANNER (Top Hero Card)
    # ----------------------------------------------------
    header_top = Inches(0.28)
    header_left = Inches(0.45)
    header_width = Inches(12.433)
    header_height = Inches(1.10)

    # Main header dark shape
    h_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, header_left, header_top, header_width, header_height)
    h_shape.fill.solid()
    h_shape.fill.fore_color.rgb = NAVY_HERO
    h_shape.line.color.rgb = MOTHERSON_BLUE
    h_shape.line.width = Pt(1.5)

    # Decorative top accent stripe on header
    stripe = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, header_left, header_top, header_width, Inches(0.06))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = MOTHERSON_BLUE
    stripe.line.fill.background()

    # Brand Emblem Badge [m]
    emblem_w = Inches(0.42)
    emblem_h = Inches(0.42)
    emblem_x = header_left + Inches(0.22)
    emblem_y = header_top + Inches(0.28)
    emblem = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emblem_x, emblem_y, emblem_w, emblem_h)
    emblem.fill.solid()
    emblem.fill.fore_color.rgb = MOTHERSON_RED
    emblem.line.fill.background()
    
    etf = emblem.text_frame
    etf.word_wrap = False
    etf.margin_left = etf.margin_top = etf.margin_right = etf.margin_bottom = 0
    ep = etf.paragraphs[0]
    ep.text = "m"
    ep.font.name = "Arial"
    ep.font.size = Pt(20)
    ep.font.bold = True
    ep.font.color.rgb = TEXT_WHITE
    ep.alignment = PP_ALIGN.CENTER

    # Brand & Title Text Box (Left side of header)
    brand_box = slide.shapes.add_textbox(header_left + Inches(0.72), header_top + Inches(0.12), Inches(6.3), Inches(0.85))
    btf = brand_box.text_frame
    btf.word_wrap = True
    btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = 0

    p_badge = btf.paragraphs[0]
    p_badge.text = "MOTHERSON GROUP  •  ENTERPRISE TALENT OS  •  v1.4.0 PRODUCTION LIVE"
    p_badge.font.name = "Arial"
    p_badge.font.size = Pt(8.5)
    p_badge.font.bold = True
    p_badge.font.color.rgb = CYAN_ACCENT
    p_badge.space_after = Pt(1)

    p_title = btf.add_paragraph()
    p_title.text = "Motherson CogniHire"
    p_title.font.name = "Arial"
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE
    p_title.space_after = Pt(1)

    p_sub = btf.add_paragraph()
    p_sub.text = "Autonomous AI Candidate Screening, Multi-Tenant Governance & Anti-Cheat Proctoring"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(9.5)
    p_sub.font.color.rgb = RGBColor(203, 213, 225)

    # 4 KPI Pills on the Right side of header
    kpis = [
        ("85% CUT", "Time-to-Shortlist", CYAN_ACCENT),
        ("99.4%", "Anti-Cheat Accuracy", EMERALD_GREEN),
        ("4-PILLAR", "Multi-Factor Scoring", AMBER_GOLD),
        ("100%", "Tenant RBAC Isolated", PURPLE_HEADER)
    ]
    pill_w = Inches(1.22)
    pill_gap = Inches(0.12)
    pill_start_x = header_left + header_width - Inches(0.2) - (4 * pill_w + 3 * pill_gap)
    pill_y = header_top + Inches(0.18)
    pill_h = Inches(0.74)

    for i, (val, label, accent_col) in enumerate(kpis):
        px = pill_start_x + i * (pill_w + pill_gap)
        p_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, pill_y, pill_w, pill_h)
        p_card.fill.solid()
        p_card.fill.fore_color.rgb = PILL_BG
        p_card.line.color.rgb = accent_col
        p_card.line.width = Pt(1.2)

        p_tb = slide.shapes.add_textbox(px, pill_y, pill_w, pill_h)
        ptf = p_tb.text_frame
        ptf.word_wrap = True
        ptf.margin_left = ptf.margin_right = ptf.margin_top = ptf.margin_bottom = Inches(0.04)
        
        p1 = ptf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        p1.text = val
        p1.font.name = "Arial"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = accent_col
        p1.space_after = Pt(1)

        p2 = ptf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = label
        p2.font.name = "Arial"
        p2.font.size = Pt(7.5)
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(241, 245, 249)

    # ----------------------------------------------------
    # MIDDLE ROW: 3 MAIN ARCHITECTURAL PILLARS
    # ----------------------------------------------------
    col_y = Inches(1.50)
    col_w = Inches(3.98)
    col_h = Inches(3.65)
    col_gap = Inches(0.24)
    c1_x = header_left
    c2_x = c1_x + col_w + col_gap
    c3_x = c2_x + col_w + col_gap

    def create_card_shell(x, y, w, h, top_border_color, header_tag, header_title):
        # Base White Card
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        # Top Accent Strip
        strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.06))
        strip.fill.solid()
        strip.fill.fore_color.rgb = top_border_color
        strip.line.fill.background()

        # Header Title Area
        tb = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.12), w - Inches(0.36), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = header_tag.upper()
        p0.font.name = "Arial"
        p0.font.size = Pt(8)
        p0.font.bold = True
        p0.font.color.rgb = top_border_color
        p0.space_after = Pt(1)

        p_t = tf.add_paragraph()
        p_t.text = header_title
        p_t.font.name = "Arial"
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_DARK

    # ====================================================
    # COLUMN 1: Multi-Tenant RBAC Hierarchy
    # ====================================================
    create_card_shell(c1_x, col_y, col_w, col_h, MOTHERSON_BLUE,
                      "01 • Multi-Tenant Architecture", "Enterprise RBAC Hierarchy")

    flow_x = c1_x + Inches(0.18)
    flow_w = col_w - Inches(0.36)

    # Tier 1: Super Admin
    t1_y = col_y + Inches(0.78)
    t1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, flow_x, t1_y, flow_w, Inches(0.48))
    t1.fill.solid()
    t1.fill.fore_color.rgb = RGBColor(239, 246, 255)
    t1.line.color.rgb = MOTHERSON_BLUE
    t1.line.width = Pt(1)

    t1_tb = slide.shapes.add_textbox(flow_x + Inches(0.08), t1_y + Inches(0.04), flow_w - Inches(0.16), Inches(0.40))
    t1_tf = t1_tb.text_frame
    t1_tf.word_wrap = True
    t1_tf.margin_left = t1_tf.margin_top = t1_tf.margin_right = t1_tf.margin_bottom = 0
    p = t1_tf.paragraphs[0]
    p.text = "👑 Motherson Super Admin (Global Owner)"
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = MOTHERSON_BLUE
    p_desc = t1_tf.add_paragraph()
    p_desc.text = "Adds enterprise clients • Creates dynamic clusters • Onboards managers"
    p_desc.font.name = "Arial"
    p_desc.font.size = Pt(7.5)
    p_desc.font.color.rgb = TEXT_MUTED

    # Tier 2: Customers
    t2_y = col_y + Inches(1.33)
    t2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, flow_x, t2_y, flow_w, Inches(0.46))
    t2.fill.solid()
    t2.fill.fore_color.rgb = RGBColor(240, 249, 255)
    t2.line.color.rgb = CYAN_ACCENT
    t2.line.width = Pt(1)

    t2_tb = slide.shapes.add_textbox(flow_x + Inches(0.08), t2_y + Inches(0.04), flow_w - Inches(0.16), Inches(0.38))
    t2_tf = t2_tb.text_frame
    t2_tf.word_wrap = True
    t2_tf.margin_left = t2_tf.margin_top = t2_tf.margin_right = t2_tf.margin_bottom = 0
    p = t2_tf.paragraphs[0]
    p.text = "🏢 Enterprise Clients (e.g., Bosch, Siemens)"
    p.font.name = "Arial"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = RGBColor(12, 74, 110)
    p_desc = t2_tf.add_paragraph()
    p_desc.text = "Complete data isolation • Dedicated customer profile & settings"
    p_desc.font.name = "Arial"
    p_desc.font.size = Pt(7.5)
    p_desc.font.color.rgb = TEXT_MUTED

    # Tier 3: Dynamic Clusters
    t3_y = col_y + Inches(1.86)
    t3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, flow_x, t3_y, flow_w, Inches(0.46))
    t3.fill.solid()
    t3.fill.fore_color.rgb = RGBColor(254, 243, 199)
    t3.line.color.rgb = AMBER_GOLD
    t3.line.width = Pt(1)

    t3_tb = slide.shapes.add_textbox(flow_x + Inches(0.08), t3_y + Inches(0.04), flow_w - Inches(0.16), Inches(0.38))
    t3_tf = t3_tb.text_frame
    t3_tf.word_wrap = True
    t3_tf.margin_left = t3_tf.margin_top = t3_tf.margin_right = t3_tf.margin_bottom = 0
    p = t3_tf.paragraphs[0]
    p.text = "⚡ Dynamic Clusters (Mechanical, SAP, Mobility...)"
    p.font.name = "Arial"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = RGBColor(146, 64, 14)
    p_desc = t3_tf.add_paragraph()
    p_desc.text = "Zero hardcoded clusters • Dynamically added per client structure"
    p_desc.font.name = "Arial"
    p_desc.font.size = Pt(7.5)
    p_desc.font.color.rgb = TEXT_MUTED

    # Tier 4: Managers & Recruiters
    t4_y = col_y + Inches(2.39)
    t4 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, flow_x, t4_y, flow_w, Inches(0.46))
    t4.fill.solid()
    t4.fill.fore_color.rgb = RGBColor(245, 243, 255)
    t4.line.color.rgb = PURPLE_ACCENT
    t4.line.width = Pt(1)

    t4_tb = slide.shapes.add_textbox(flow_x + Inches(0.08), t4_y + Inches(0.04), flow_w - Inches(0.16), Inches(0.38))
    t4_tf = t4_tb.text_frame
    t4_tf.word_wrap = True
    t4_tf.margin_left = t4_tf.margin_top = t4_tf.margin_right = t4_tf.margin_bottom = 0
    p = t4_tf.paragraphs[0]
    p.text = "👔 Cluster Managers & 🎯 Active Recruiters"
    p.font.name = "Arial"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = RGBColor(109, 40, 217)
    p_desc = t4_tf.add_paragraph()
    p_desc.text = "Managers post/close jobs • Recruiters upload CVs & execute screening"
    p_desc.font.name = "Arial"
    p_desc.font.size = Pt(7.5)
    p_desc.font.color.rgb = TEXT_MUTED

    # Col 1 Key Capabilities Bullet Box
    b1_y = col_y + Inches(2.93)
    b1_tb = slide.shapes.add_textbox(flow_x, b1_y, flow_w, Inches(0.65))
    b1_tf = b1_tb.text_frame
    b1_tf.word_wrap = True
    b1_tf.margin_left = b1_tf.margin_top = b1_tf.margin_right = b1_tf.margin_bottom = 0

    points_1 = [
        "Dedicated Credentials: Independent secure login per role & entity.",
        "Role Switcher: 1-click persona switching for effortless stakeholder demo.",
        "Role Dashboards: Scoped KPIs, analytics & Recharts visualizations."
    ]
    for pt in points_1:
        p = b1_tf.add_paragraph() if b1_tf.paragraphs[0].text else b1_tf.paragraphs[0]
        p.text = "✔ " + pt
        p.font.name = "Arial"
        p.font.size = Pt(7.8)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(2)

    # ====================================================
    # COLUMN 2: 4-Pillar AI Fit Scoring & Radar Matrix
    # ====================================================
    create_card_shell(c2_x, col_y, col_w, col_h, RGBColor(109, 40, 217),
                      "02 • Multi-Factor Intelligence", "4-Pillar AI Fit Scoring")

    pbox_x = c2_x + Inches(0.18)
    pbox_w = col_w - Inches(0.36)

    pillars = [
        ("Technical Skills", "40%", "Hard requirements match, must-have tools & depth", RGBColor(238, 242, 255), RGBColor(79, 70, 229)),
        ("Experience Fit", "25%", "Tenure progression, seniority, leadership & impact", RGBColor(240, 249, 255), CYAN_ACCENT),
        ("Domain Relevance", "20%", "Automotive, industrial, manufacturing industry alignment", RGBColor(236, 253, 245), EMERALD_GREEN),
        ("Education & Creds", "15%", "Degree relevance, tier-1 institutions & certifications", RGBColor(254, 243, 199), AMBER_GOLD)
    ]

    p_item_y = col_y + Inches(0.78)
    item_h = Inches(0.48)
    gap_item = Inches(0.06)

    for idx, (p_title, p_pct, p_desc, p_bg, p_col) in enumerate(pillars):
        curr_y = p_item_y + idx * (item_h + gap_item)
        
        # Pillar mini box
        p_sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pbox_x, curr_y, pbox_w, item_h)
        p_sh.fill.solid()
        p_sh.fill.fore_color.rgb = p_bg
        p_sh.line.color.rgb = p_col
        p_sh.line.width = Pt(1)

        p_tb = slide.shapes.add_textbox(pbox_x + Inches(0.08), curr_y + Inches(0.04), pbox_w - Inches(0.16), item_h - Inches(0.08))
        ptf = p_tb.text_frame
        ptf.word_wrap = True
        ptf.margin_left = ptf.margin_top = ptf.margin_right = ptf.margin_bottom = 0

        p = ptf.paragraphs[0]
        p.text = f"{p_title}  [{p_pct} Weight]"
        p.font.name = "Arial"
        p.font.size = Pt(8.8)
        p.font.bold = True
        p.font.color.rgb = p_col

        p2 = ptf.add_paragraph()
        p2.text = p_desc
        p2.font.name = "Arial"
        p2.font.size = Pt(7.5)
        p2.font.color.rgb = TEXT_DARK

    # Comparative Intelligence Box
    comp_y = col_y + Inches(2.93)
    comp_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pbox_x, comp_y, pbox_w, Inches(0.62))
    comp_box.fill.solid()
    comp_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
    comp_box.line.color.rgb = CARD_BORDER
    comp_box.line.width = Pt(1)

    c_tb = slide.shapes.add_textbox(pbox_x + Inches(0.08), comp_y + Inches(0.05), pbox_w - Inches(0.16), Inches(0.52))
    ctf = c_tb.text_frame
    ctf.word_wrap = True
    ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0

    p = ctf.paragraphs[0]
    p.text = "⚖️ Head-to-Head Comparative Intelligence"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(109, 40, 217)

    p2 = ctf.add_paragraph()
    p2.text = "• Automated Top 5 Radar delta breakdown (Rank #1 vs Runner-up).\n• Claude & Gemini synthesis memos eliminate keyword-stuffing bias."
    p2.font.name = "Arial"
    p2.font.size = Pt(7.3)
    p2.font.color.rgb = TEXT_DARK

    # ====================================================
    # COLUMN 3: Zero-Trust Anti-Cheat & Video Proctoring
    # ====================================================
    create_card_shell(c3_x, col_y, col_w, col_h, ROSE_ACCENT,
                      "03 • Zero-Trust Security", "AI Anti-Cheat & Proctoring")

    c3box_x = c3_x + Inches(0.18)
    c3box_w = col_w - Inches(0.36)

    cheats = [
        ("🛡️ Prompt-Injection Defense & 5-Gram Echo", "Strips hidden white-text instructions; penalizes verbatim JD sentence copying.", ROSE_ACCENT),
        ("👁️ Sustained Gaze & Teleprompter Tracking", "Flags candidate downward glances (>3.5s) at smartphone screens / notes.", CYAN_ACCENT),
        ("⌨️ Keystroke Cadence Dynamics", "Detects instant clipboard dumps & superhuman WPM bursts lacking thought pauses.", EMERALD_GREEN),
        ("🤖 ChatGPT Stylometry Audit", "Calculates AI Content Probability % and flags signature synthetic bot phrasing.", PURPLE_ACCENT)
    ]

    c3_item_y = col_y + Inches(0.78)
    for idx, (title, desc, color) in enumerate(cheats):
        curr_y = c3_item_y + idx * (item_h + gap_item)

        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c3box_x, curr_y, c3box_w, item_h)
        sh.fill.solid()
        sh.fill.fore_color.rgb = ROW_BG_LIGHT
        sh.line.color.rgb = CARD_BORDER
        sh.line.width = Pt(1)

        # left mini color bar
        cbar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c3box_x, curr_y, Inches(0.06), item_h)
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = color
        cbar.line.fill.background()

        tb = slide.shapes.add_textbox(c3box_x + Inches(0.12), curr_y + Inches(0.04), c3box_w - Inches(0.2), item_h - Inches(0.08))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Arial"
        p2.font.size = Pt(7.3)
        p2.font.color.rgb = TEXT_MUTED

    # Proctored Timer Control Box
    timer_y = col_y + Inches(2.93)
    timer_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c3box_x, timer_y, c3box_w, Inches(0.62))
    timer_box.fill.solid()
    timer_box.fill.fore_color.rgb = RGBColor(255, 241, 242)
    timer_box.line.color.rgb = RGBColor(254, 205, 211)
    timer_box.line.width = Pt(1)

    t_tb = slide.shapes.add_textbox(c3box_x + Inches(0.08), timer_y + Inches(0.05), c3box_w - Inches(0.16), Inches(0.52))
    ttf = t_tb.text_frame
    ttf.word_wrap = True
    ttf.margin_left = ttf.margin_top = ttf.margin_right = ttf.margin_bottom = 0

    p = ttf.paragraphs[0]
    p.text = "⏱️ Proctored Environment & Global Countdown"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = ROSE_ACCENT

    p2 = ttf.add_paragraph()
    p2.text = "• Strict 30-min global exam clock with live MediaPipe camera telemetry.\n• Real-time behavioral anomaly logs generated for executive audit."
    p2.font.name = "Arial"
    p2.font.size = Pt(7.3)
    p2.font.color.rgb = TEXT_DARK

    # ----------------------------------------------------
    # BOTTOM ROW: 3 CAPABILITY CARDS
    # ----------------------------------------------------
    b_row_y = Inches(5.24)
    b_row_h = Inches(1.72)
    b_col_w = Inches(3.98)

    # CARD 4: Job Lifecycle & Client Manager
    create_card_shell(c1_x, b_row_y, b_col_w, b_row_h, EMERALD_GREEN,
                      "04 • Job Requisition Control", "Job Lifecycle & Client Manager")

    b4_tb = slide.shapes.add_textbox(c1_x + Inches(0.18), b_row_y + Inches(0.72), b_col_w - Inches(0.36), Inches(0.92))
    b4_tf = b4_tb.text_frame
    b4_tf.word_wrap = True
    b4_tf.margin_left = b4_tf.margin_top = b4_tf.margin_right = b4_tf.margin_bottom = 0

    b4_points = [
        "3-Stage Pipeline: Draft Requisition ➔ [Post to Recruiters] ➔ [Close Job].",
        "Client Manager Name: Directly assigned in JD (e.g., Klaus Richter - Bosch).",
        "Stage Gates: Recruiter upload unlocked only when manager opens requisition."
    ]
    for pt in b4_points:
        p = b4_tf.add_paragraph() if b4_tf.paragraphs[0].text else b4_tf.paragraphs[0]
        p.text = "• " + pt
        p.font.name = "Arial"
        p.font.size = Pt(8)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(2)

    # CARD 5: Reversible HR Decisions
    create_card_shell(c2_x, b_row_y, b_col_w, b_row_h, CYAN_ACCENT,
                      "05 • Decision Governance", "Reversible Hiring Decisions")

    b5_tb = slide.shapes.add_textbox(c2_x + Inches(0.18), b_row_y + Inches(0.72), b_col_w - Inches(0.36), Inches(0.92))
    b5_tf = b5_tb.text_frame
    b5_tf.word_wrap = True
    b5_tf.margin_left = b5_tf.margin_top = b5_tf.margin_right = b5_tf.margin_bottom = 0

    b5_points = [
        "Dynamic Decision Toggle: 1-click switch [✓ Selected] / [✕ Not Selected].",
        "Live Vacancy Counter: Real-time progress bar (e.g. 3 of 4 Positions Staffed).",
        "Zero-Risk Undo: Instant reversal & reset safety to prevent HR errors."
    ]
    for pt in b5_points:
        p = b5_tf.add_paragraph() if b5_tf.paragraphs[0].text else b5_tf.paragraphs[0]
        p.text = "• " + pt
        p.font.name = "Arial"
        p.font.size = Pt(8)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(2)

    # CARD 6: Multi-Sheet Excel Reports & Cloud Sync
    create_card_shell(c3_x, b_row_y, b_col_w, b_row_h, AMBER_GOLD,
                      "06 • Reporting & Persistence", "Multi-Sheet Excel & Cloud Sync")

    b6_tb = slide.shapes.add_textbox(c3_x + Inches(0.18), b_row_y + Inches(0.72), b_col_w - Inches(0.36), Inches(0.92))
    b6_tf = b6_tb.text_frame
    b6_tf.word_wrap = True
    b6_tf.margin_left = b6_tf.margin_top = b6_tf.margin_right = b6_tf.margin_bottom = 0

    b6_points = [
        "Multi-Sheet Excel (.xlsx): Sheet 1 (Openings/Pipeline) + Sheet 2 (Candidates/Scores).",
        "1-Click Export: In-browser SheetJS generation with zero backend latency.",
        "Dual Persistence: Local ACID SQLite (sift.db) + Cloud auto-sync (GCS)."
    ]
    for pt in b6_points:
        p = b6_tf.add_paragraph() if b6_tf.paragraphs[0].text else b6_tf.paragraphs[0]
        p.text = "• " + pt
        p.font.name = "Arial"
        p.font.size = Pt(8)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(2)

    # ----------------------------------------------------
    # FOOTER BAR
    # ----------------------------------------------------
    foot_y = Inches(7.06)
    foot_h = Inches(0.28)
    foot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, header_left, foot_y, header_width, foot_h)
    foot.fill.solid()
    foot.fill.fore_color.rgb = NAVY_HERO
    foot.line.fill.background()

    foot_tb = slide.shapes.add_textbox(header_left + Inches(0.2), foot_y + Inches(0.04), header_width - Inches(0.4), foot_h - Inches(0.08))
    ftf = foot_tb.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0

    fp = ftf.paragraphs[0]
    fp.text = "Motherson CogniHire v1.4.0  •  Enterprise Autonomous Candidate Screening & Proctoring System  •  Confidential & Proprietary  •  © 2026 Motherson Group"
    fp.font.name = "Arial"
    fp.font.size = Pt(8)
    fp.font.color.rgb = RGBColor(148, 163, 184)
    fp.alignment = PP_ALIGN.CENTER

    prs.save(output_path)
    print(f"Successfully generated refined single-page presentation: {output_path}")

if __name__ == "__main__":
    out_dir = os.path.dirname(os.path.abspath(__file__))
    target_pptx = os.path.join(out_dir, "Motherson_CogniHire_Executive_Overview.pptx")
    build_single_page_presentation(target_pptx)
