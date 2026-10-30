import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_single_page_presentation_with_pictures(output_path, assets_dir):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Colors
    BG_CANVAS        = RGBColor(241, 245, 249)
    NAVY_HERO        = RGBColor(11, 23, 44)
    MOTHERSON_BLUE   = RGBColor(3, 77, 161)
    MOTHERSON_RED    = RGBColor(225, 29, 72)
    CYAN_ACCENT      = RGBColor(14, 165, 233)
    EMERALD_GREEN    = RGBColor(16, 185, 129)
    AMBER_GOLD       = RGBColor(245, 158, 11)
    PURPLE_ACCENT    = RGBColor(124, 58, 237)
    PURPLE_HEADER    = RGBColor(192, 132, 252)
    ROSE_ACCENT      = RGBColor(244, 63, 94)
    CARD_BG          = RGBColor(255, 255, 255)
    CARD_BORDER      = RGBColor(203, 213, 225)
    PILL_BG          = RGBColor(23, 37, 63)
    TEXT_DARK        = RGBColor(15, 23, 42)
    TEXT_MUTED       = RGBColor(100, 116, 139)
    TEXT_WHITE       = RGBColor(255, 255, 255)
    ROW_BG_LIGHT     = RGBColor(248, 250, 252)

    # 1. Canvas Background
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BG_CANVAS
    bg_shape.line.fill.background()

    # ----------------------------------------------------
    # HEADER BANNER (Top Hero Card)
    # ----------------------------------------------------
    header_top = Inches(0.24)
    header_left = Inches(0.45)
    header_width = Inches(12.433)
    header_height = Inches(1.10)

    h_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, header_left, header_top, header_width, header_height)
    h_shape.fill.solid()
    h_shape.fill.fore_color.rgb = NAVY_HERO
    h_shape.line.color.rgb = MOTHERSON_BLUE
    h_shape.line.width = Pt(1.5)

    stripe = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, header_left, header_top, header_width, Inches(0.06))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = MOTHERSON_BLUE
    stripe.line.fill.background()

    # Brand Emblem [m]
    emblem = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, header_left + Inches(0.20), header_top + Inches(0.26), Inches(0.42), Inches(0.42))
    emblem.fill.solid()
    emblem.fill.fore_color.rgb = MOTHERSON_RED
    emblem.line.fill.background()
    etf = emblem.text_frame
    etf.word_wrap = False
    etf.margin_left = etf.margin_top = etf.margin_right = etf.margin_bottom = 0
    ep = etf.paragraphs[0]
    ep.text = "m"
    ep.font.name = "Arial"
    ep.font.size = Pt(20)
    ep.font.bold = True
    ep.font.color.rgb = TEXT_WHITE
    ep.alignment = PP_ALIGN.CENTER

    # Brand Title
    brand_box = slide.shapes.add_textbox(header_left + Inches(0.70), header_top + Inches(0.12), Inches(6.3), Inches(0.85))
    btf = brand_box.text_frame
    btf.word_wrap = True
    btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = 0

    p_badge = btf.paragraphs[0]
    p_badge.text = "MOTHERSON GROUP  •  ENTERPRISE TALENT OS  •  v1.4.0 PRODUCTION LIVE"
    p_badge.font.name = "Arial"
    p_badge.font.size = Pt(8.5)
    p_badge.font.bold = True
    p_badge.font.color.rgb = CYAN_ACCENT
    p_badge.space_after = Pt(1)

    p_title = btf.add_paragraph()
    p_title.text = "Motherson CogniHire"
    p_title.font.name = "Arial"
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE
    p_title.space_after = Pt(1)

    p_sub = btf.add_paragraph()
    p_sub.text = "Autonomous AI Candidate Screening, Multi-Tenant Governance & Anti-Cheat Proctoring"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(9.5)
    p_sub.font.color.rgb = RGBColor(203, 213, 225)

    # 4 KPI Pills
    kpis = [
        ("85% CUT", "Time-to-Shortlist", CYAN_ACCENT),
        ("99.4%", "Anti-Cheat Accuracy", EMERALD_GREEN),
        ("4-PILLAR", "Multi-Factor Scoring", AMBER_GOLD),
        ("100%", "Tenant RBAC Isolated", PURPLE_HEADER)
    ]
    pill_w = Inches(1.22)
    pill_gap = Inches(0.12)
    pill_start_x = header_left + header_width - Inches(0.2) - (4 * pill_w + 3 * pill_gap)
    pill_y = header_top + Inches(0.18)
    pill_h = Inches(0.74)

    for i, (val, label, accent_col) in enumerate(kpis):
        px = pill_start_x + i * (pill_w + pill_gap)
        p_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, pill_y, pill_w, pill_h)
        p_card.fill.solid()
        p_card.fill.fore_color.rgb = PILL_BG
        p_card.line.color.rgb = accent_col
        p_card.line.width = Pt(1.2)

        p_tb = slide.shapes.add_textbox(px, pill_y, pill_w, pill_h)
        ptf = p_tb.text_frame
        ptf.word_wrap = True
        ptf.margin_left = ptf.margin_right = ptf.margin_top = ptf.margin_bottom = Inches(0.04)
        
        p1 = ptf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        p1.text = val
        p1.font.name = "Arial"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = accent_col
        p1.space_after = Pt(1)

        p2 = ptf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = label
        p2.font.name = "Arial"
        p2.font.size = Pt(7.5)
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(241, 245, 249)

    # ----------------------------------------------------
    # MIDDLE ROW: 3 MAIN COLUMNS WITH EMBEDDED SCREENSHOTS
    # ----------------------------------------------------
    col_y = Inches(1.44)
    col_w = Inches(3.98)
    col_h = Inches(3.75)
    col_gap = Inches(0.24)
    c1_x = header_left
    c2_x = c1_x + col_w + col_gap
    c3_x = c2_x + col_w + col_gap

    def create_card_shell(x, y, w, h, top_border_color, header_tag, header_title):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.06))
        strip.fill.solid()
        strip.fill.fore_color.rgb = top_border_color
        strip.line.fill.background()

        tb = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.10), w - Inches(0.36), Inches(0.55))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = header_tag.upper()
        p0.font.name = "Arial"
        p0.font.size = Pt(7.8)
        p0.font.bold = True
        p0.font.color.rgb = top_border_color
        p0.space_after = Pt(1)

        p_t = tf.add_paragraph()
        p_t.text = header_title
        p_t.font.name = "Arial"
        p_t.font.size = Pt(12.5)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_DARK

    # ====================================================
    # COLUMN 1: Multi-Tenant RBAC Hierarchy + Screenshot
    # ====================================================
    create_card_shell(c1_x, col_y, col_w, col_h, MOTHERSON_BLUE,
                      "01 • Hierarchy & Governance", "Enterprise RBAC Hierarchy")

    flow_x = c1_x + Inches(0.18)
    flow_w = col_w - Inches(0.36)

    # Bullet list of hierarchy
    c1_tb = slide.shapes.add_textbox(flow_x, col_y + Inches(0.70), flow_w, Inches(1.50))
    c1_tf = c1_tb.text_frame
    c1_tf.word_wrap = True
    c1_tf.margin_left = c1_tf.margin_top = c1_tf.margin_right = c1_tf.margin_bottom = 0

    c1_points = [
        ("👑 Motherson Super Admin", "Global platform owner; configures enterprise clients & clusters."),
        ("🏢 Enterprise Customers", "Isolated tenants (Bosch, Siemens) with dynamic scoping."),
        ("⚡ Dynamic Clusters", "Zero hardcoded clusters; created on-the-fly (Mechanical, SAP)."),
        ("👔 Managers & 🎯 Recruiters", "Cluster managers post jobs; recruiters screen candidates.")
    ]
    for idx, (title, desc) in enumerate(c1_points):
        p = c1_tf.add_paragraph() if idx > 0 else c1_tf.paragraphs[0]
        p.text = f"{title}: "
        p.font.name = "Arial"
        p.font.size = Pt(8.2)
        p.font.bold = True
        p.font.color.rgb = MOTHERSON_BLUE
        
        # description
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_DARK
        p.space_after = Pt(3)

    # EMBED SCREENSHOT 1: Admin Command Center picture
    admin_pic_path = os.path.join(assets_dir, "asset_admin.png")
    if os.path.exists(admin_pic_path):
        pic_y = col_y + Inches(2.28)
        pic_w = flow_w
        pic_h = Inches(1.28)
        slide.shapes.add_picture(admin_pic_path, flow_x, pic_y, width=pic_w, height=pic_h)

        # Micro caption
        cap_tb = slide.shapes.add_textbox(flow_x, pic_y + pic_h + Inches(0.02), flow_w, Inches(0.16))
        ctf = cap_tb.text_frame
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0
        cp = ctf.paragraphs[0]
        cp.text = "Live UI: Motherson Super Admin Command Center & Dynamic Clusters"
        cp.font.name = "Arial"
        cp.font.size = Pt(6.8)
        cp.font.italic = True
        cp.font.color.rgb = TEXT_MUTED

    # ====================================================
    # COLUMN 2: 4-Pillar AI Fit Scoring + Candidate Card Screenshot
    # ====================================================
    create_card_shell(c2_x, col_y, col_w, col_h, PURPLE_ACCENT,
                      "02 • Multi-Factor Intelligence", "4-Pillar AI Fit Scoring")

    pbox_x = c2_x + Inches(0.18)
    pbox_w = col_w - Inches(0.36)

    pillars = [
        ("Technical Skills (40%)", "Hard requirements match & tool proficiency", PURPLE_ACCENT),
        ("Experience Fit (25%)", "Tenure progression, seniority & impact", CYAN_ACCENT),
        ("Domain Relevance (20%)", "Automotive & industrial engineering depth", EMERALD_GREEN),
        ("Education & Creds (15%)", "Degree pedigree, verified credentials", AMBER_GOLD)
    ]
    p_item_y = col_y + Inches(0.70)
    p_h = Inches(0.36)
    for idx, (p_title, p_desc, p_col) in enumerate(pillars):
        curr_y = p_item_y + idx * Inches(0.38)
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pbox_x, curr_y, pbox_w, p_h)
        sh.fill.solid()
        sh.fill.fore_color.rgb = ROW_BG_LIGHT
        sh.line.color.rgb = p_col
        sh.line.width = Pt(1)

        # side bar
        cbar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pbox_x, curr_y, Inches(0.06), p_h)
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = p_col
        cbar.line.fill.background()

        tb = slide.shapes.add_textbox(pbox_x + Inches(0.12), curr_y + Inches(0.02), pbox_w - Inches(0.18), p_h - Inches(0.04))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"{p_title} — "
        p.font.name = "Arial"
        p.font.size = Pt(8.2)
        p.font.bold = True
        p.font.color.rgb = p_col

        run = p.add_run()
        run.text = p_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_DARK

    # EMBED SCREENSHOT 2: Real Shortlist Candidate Card (88% Fit, 100% Authentic)
    cand_pic_path = os.path.join(assets_dir, "asset_candidate.png")
    if os.path.exists(cand_pic_path):
        pic_y = col_y + Inches(2.28)
        pic_w = flow_w
        pic_h = Inches(1.28)
        slide.shapes.add_picture(cand_pic_path, pbox_x, pic_y, width=pic_w, height=pic_h)

        cap_tb = slide.shapes.add_textbox(pbox_x, pic_y + pic_h + Inches(0.02), flow_w, Inches(0.16))
        ctf = cap_tb.text_frame
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0
        cp = ctf.paragraphs[0]
        cp.text = "Live UI: Candidate Scorecard with 88% Match & 100% Authenticity"
        cp.font.name = "Arial"
        cp.font.size = Pt(6.8)
        cp.font.italic = True
        cp.font.color.rgb = TEXT_MUTED

    # ====================================================
    # COLUMN 3: Zero-Trust Anti-Cheat + Proctoring Screenshot
    # ====================================================
    create_card_shell(c3_x, col_y, col_w, col_h, ROSE_ACCENT,
                      "03 • Zero-Trust Security", "AI Anti-Cheat & Proctoring")

    c3box_x = c3_x + Inches(0.18)
    c3box_w = col_w - Inches(0.36)

    c3_tb = slide.shapes.add_textbox(c3box_x, col_y + Inches(0.70), c3box_w, Inches(1.50))
    c3_tf = c3_tb.text_frame
    c3_tf.word_wrap = True
    c3_tf.margin_left = c3_tf.margin_top = c3_tf.margin_right = c3_tf.margin_bottom = 0

    cheats = [
        ("🛡️ Prompt-Injection Defense", "Strips hidden text; penalizes verbatim JD echo."),
        ("👁️ Gaze & Teleprompter", "Flags downward glances (>3.5s) at smartphones."),
        ("⌨️ Keystroke Cadence Dynamics", "Detects clipboard dumps & high WPM typing bursts."),
        ("🤖 ChatGPT Stylometry Audit", "Calculates AI Probability % & synthetic bot phrasing.")
    ]
    for idx, (title, desc) in enumerate(cheats):
        p = c3_tf.add_paragraph() if idx > 0 else c3_tf.paragraphs[0]
        p.text = f"{title}: "
        p.font.name = "Arial"
        p.font.size = Pt(8.2)
        p.font.bold = True
        p.font.color.rgb = ROSE_ACCENT

        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_DARK
        p.space_after = Pt(3)

    # EMBED SCREENSHOT 3: Live Proctored Interview Portal picture
    proc_pic_path = os.path.join(assets_dir, "asset_proctor.png")
    if os.path.exists(proc_pic_path):
        pic_y = col_y + Inches(2.28)
        # centered thumbnail
        pic_w = Inches(2.20)
        pic_h = Inches(1.28)
        pic_x = c3box_x + (c3box_w - pic_w) / 2
        slide.shapes.add_picture(proc_pic_path, pic_x, pic_y, width=pic_w, height=pic_h)

        cap_tb = slide.shapes.add_textbox(c3box_x, pic_y + pic_h + Inches(0.02), c3box_w, Inches(0.16))
        ctf = cap_tb.text_frame
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        cp.text = "Live UI: 30-Min Proctored Assessment with Gaze Malpractice Rules"
        cp.font.name = "Arial"
        cp.font.size = Pt(6.8)
        cp.font.italic = True
        cp.font.color.rgb = TEXT_MUTED

    # ----------------------------------------------------
    # BOTTOM ROW: 3 OPERATIONAL CAPABILITY CARDS
    # ----------------------------------------------------
    b_row_y = Inches(5.34)
    b_row_h = Inches(1.64)
    b_col_w = Inches(3.98)

    # CARD 4: Job Lifecycle & Client Manager
    create_card_shell(c1_x, b_row_y, b_col_w, b_row_h, EMERALD_GREEN,
                      "04 • Job Requisition Control", "Job Lifecycle & Client Manager")

    b4_tb = slide.shapes.add_textbox(c1_x + Inches(0.18), b_row_y + Inches(0.68), b_col_w - Inches(0.36), Inches(0.88))
    b4_tf = b4_tb.text_frame
    b4_tf.word_wrap = True
    b4_tf.margin_left = b4_tf.margin_top = b4_tf.margin_right = b4_tf.margin_bottom = 0

    b4_points = [
        "3-Stage Pipeline: Draft Requisition ➔ [Post to Recruiters] ➔ [Close Job].",
        "Client Manager Name: Directly assigned in JD (e.g., Klaus Richter - Bosch).",
        "Stage Gates: Recruiter upload unlocked only when manager opens requisition."
    ]
    for pt in b4_points:
        p = b4_tf.add_paragraph() if b4_tf.paragraphs[0].text else b4_tf.paragraphs[0]
        p.text = "• " + pt
        p.font.name = "Arial"
        p.font.size = Pt(7.8)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(2)

    # CARD 5: Reversible HR Decisions
    create_card_shell(c2_x, b_row_y, b_col_w, b_row_h, CYAN_ACCENT,
                      "05 • Decision Governance", "Reversible Hiring Decisions")

    b5_tb = slide.shapes.add_textbox(c2_x + Inches(0.18), b_row_y + Inches(0.68), b_col_w - Inches(0.36), Inches(0.88))
    b5_tf = b5_tb.text_frame
    b5_tf.word_wrap = True
    b5_tf.margin_left = b5_tf.margin_top = b5_tf.margin_right = b5_tf.margin_bottom = 0

    b5_points = [
        "Dynamic Decision Toggle: 1-click switch [✓ Selected] / [✕ Not Selected].",
        "Live Vacancy Counter: Real-time progress bar (e.g. 3 of 4 Positions Staffed).",
        "Zero-Risk Undo: Instant reversal & reset safety to prevent HR errors."
    ]
    for pt in b5_points:
        p = b5_tf.add_paragraph() if b5_tf.paragraphs[0].text else b5_tf.paragraphs[0]
        p.text = "• " + pt
        p.font.name = "Arial"
        p.font.size = Pt(7.8)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(2)

    # CARD 6: Multi-Sheet Excel Reports & Cloud Sync
    create_card_shell(c3_x, b_row_y, b_col_w, b_row_h, AMBER_GOLD,
                      "06 • Reporting & Persistence", "Multi-Sheet Excel & Cloud Sync")

    b6_tb = slide.shapes.add_textbox(c3_x + Inches(0.18), b_row_y + Inches(0.68), b_col_w - Inches(0.36), Inches(0.88))
    b6_tf = b6_tb.text_frame
    b6_tf.word_wrap = True
    b6_tf.margin_left = b6_tf.margin_top = b6_tf.margin_right = b6_tf.margin_bottom = 0

    b6_points = [
        "Multi-Sheet Excel (.xlsx): Sheet 1 (Openings/Pipeline) + Sheet 2 (Candidates/Scores).",
        "1-Click Export: In-browser SheetJS generation with zero backend latency.",
        "Dual Persistence: Local ACID SQLite (sift.db) + Cloud auto-sync (GCS)."
    ]
    for pt in b6_points:
        p = b6_tf.add_paragraph() if b6_tf.paragraphs[0].text else b6_tf.paragraphs[0]
        p.text = "• " + pt
        p.font.name = "Arial"
        p.font.size = Pt(7.8)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(2)

    # ----------------------------------------------------
    # FOOTER BAR
    # ----------------------------------------------------
    foot_y = Inches(7.08)
    foot_h = Inches(0.26)
    foot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, header_left, foot_y, header_width, foot_h)
    foot.fill.solid()
    foot.fill.fore_color.rgb = NAVY_HERO
    foot.line.fill.background()

    foot_tb = slide.shapes.add_textbox(header_left + Inches(0.2), foot_y + Inches(0.03), header_width - Inches(0.4), foot_h - Inches(0.06))
    ftf = foot_tb.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0

    fp = ftf.paragraphs[0]
    fp.text = "Motherson CogniHire v1.4.0  •  Enterprise Autonomous Candidate Screening & Proctoring System  •  Confidential & Proprietary  •  © 2026 Motherson Group"
    fp.font.name = "Arial"
    fp.font.size = Pt(7.8)
    fp.font.color.rgb = RGBColor(148, 163, 184)
    fp.alignment = PP_ALIGN.CENTER

    prs.save(output_path)
    print(f"Successfully generated single-page presentation with screenshots: {output_path}")

if __name__ == "__main__":
    assets_dir = r"D:\AI_Resume_Shortlisting"
    out_dir = r"D:\AI_Resume_Shortlisting\sift-resume-agent\sift-resume-agent"
    target_pptx = os.path.join(out_dir, "Motherson_CogniHire_Executive_Overview.pptx")
    build_single_page_presentation_with_pictures(target_pptx, assets_dir)
